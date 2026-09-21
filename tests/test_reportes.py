from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from pydantic import ValidationError

from app.modules.inteligencia.exports.reporte_excel import crear_excel
from app.modules.inteligencia.exports.reporte_pdf import crear_pdf
from app.modules.inteligencia.schemas.reportes.reporte_request import ReporteRequest
from app.modules.inteligencia.services import reporte_service


def test_rango_invertido_no_se_acepta():
    try:
        ReporteRequest(tipo="VENTAS", fecha_desde=date(2026, 9, 15), fecha_hasta=date(2026, 9, 1))
    except ValidationError:
        return
    raise AssertionError("El rango invertido debió rechazarse")


def test_voz_interpreta_inventario_sin_confundirlo_con_ventas(monkeypatch):
    monkeypatch.setattr(reporte_service, "listar_sucursales", lambda: [{"id": 1, "nombre": "Sucursal Centro"}])
    respuesta = reporte_service.interpretar("Inventario bajo stock en Sucursal Centro")
    assert respuesta["interpretado"] is True
    assert respuesta["filtros"]["tipo"] == "INVENTARIO"
    assert respuesta["filtros"]["solo_bajo_stock"] is True
    assert respuesta["filtros"]["sucursal_id"] == 1


def test_ventas_y_exportaciones_comparten_los_totales(monkeypatch):
    monkeypatch.setattr(reporte_service, "listar_sucursales", lambda: [{"id": 1, "nombre": "Sucursal Centro"}])
    fecha = datetime(2026, 9, 10, 12, tzinfo=ZoneInfo("America/La_Paz"))
    monkeypatch.setattr(reporte_service, "consultar_reporte", lambda _: [
        {"codigo": "V1", "fecha": fecha, "sucursal": "Sucursal Centro", "subtotal": Decimal("40.00"), "descuento": Decimal("0.00"), "total": Decimal("40.00"), "unidades": 2},
        {"codigo": "V2", "fecha": fecha, "sucursal": "Sucursal Centro", "subtotal": Decimal("30.00"), "descuento": Decimal("0.00"), "total": Decimal("30.00"), "unidades": 1},
    ])
    reporte = reporte_service.generar_reporte(ReporteRequest(tipo="VENTAS", fecha_desde=date(2026, 9, 1), fecha_hasta=date(2026, 9, 15)))
    assert reporte["total_filas"] == 2
    assert reporte["indicadores"][2]["valor"] == 70.0
    assert reporte["indicadores"][1]["valor"] == 3
    assert crear_pdf(reporte).startswith(b"%PDF-")
    libro = load_workbook(BytesIO(crear_excel(reporte)))
    assert libro.sheetnames == ["Resumen", "Datos"]
    assert any(row[0].value == "Total ventas (Bs)" and row[1].value == 70.0 for row in libro["Resumen"].iter_rows(min_col=1, max_col=2))
    assert libro["Datos"].max_row == 3
