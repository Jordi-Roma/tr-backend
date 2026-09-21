from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def _seguro(valor: object) -> object:
    if isinstance(valor, str) and valor.startswith(("=", "+", "-", "@")):
        return "'" + valor
    return valor


def crear_excel(reporte: dict[str, object]) -> bytes:
    wb = Workbook()
    resumen = wb.active
    resumen.title = "Resumen"
    resumen.append(["StyleAR", reporte["titulo"]])
    resumen.append(["Generado", reporte["generado_en"]])
    if reporte.get("nota"):
        resumen.append(["Nota", reporte["nota"]])
    for key, value in reporte["filtros_efectivos"].items():
        resumen.append([key.replace("_", " ").title(), str(value) if value is not None else "Todas"])
    resumen.append([])
    for indicador in reporte["indicadores"]:
        resumen.append([indicador["label"], indicador["valor"]])
    datos = wb.create_sheet("Datos")
    columnas = reporte["columnas"]
    if columnas:
        datos.append([c["label"] for c in columnas])
        for fila in reporte["filas"]:
            datos.append([_seguro(fila.get(c["key"], "")) for c in columnas])
        datos.auto_filter.ref = datos.dimensions
        datos.freeze_panes = "A2"
    else:
        datos.append(["Sin datos para los filtros seleccionados"])
    for sheet in (resumen, datos):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="263043")
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(42, max(14, max(len(str(c.value or "")) for c in column) + 2))
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
